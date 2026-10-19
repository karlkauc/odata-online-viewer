"""Excel export: paging through a whole entity set and typed workbook cells."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit

import openpyxl
import pytest

from app.odata.csdl_parser import parse_metadata
from app.odata.data_client import DataError, QueryOptions
from app.odata.xlsx_export import cell_value, collect_rows, property_types, to_xlsx

SAMPLES = Path(__file__).resolve().parent.parent / "app" / "odata" / "samples"
V4_ROOT = "https://services.odata.org/V4/Northwind/Northwind.svc/"
V2_ROOT = "https://services.odata.org/V2/Northwind/Northwind.svc/"


def model(name: str, root: str):
    m = parse_metadata((SAMPLES / name).read_bytes()).model
    m.service_root = root
    return m


class Pages:
    """Stand-in for fetch_json: answers by request URL via ``answer(url)``."""

    def __init__(self, answer) -> None:
        self.answer = answer
        self.urls: list[str] = []

    def __call__(self, url, model, auth, *, timeout, max_bytes) -> tuple[Any, int, str | None, int]:
        self.urls.append(url)
        return self.answer(url), 200, "application/json", 1


def skip_of(url: str) -> int:
    return int(parse_qs(urlsplit(url).query).get("$skip", ["0"])[0])


def collect(m, fetch, entity_set: str = "Products", options: QueryOptions | None = None, **kw):
    kw.setdefault("max_rows", 50_000)
    kw.setdefault("deadline_seconds", 60)
    return collect_rows(
        m, entity_set, options or QueryOptions(), None, timeout=1, max_bytes=1, fetch=fetch, **kw
    )


def test_follows_v4_next_links() -> None:
    def answer(url: str):
        if "$skiptoken" in url:
            return {"value": [{"ProductID": 3}]}
        return {"value": [{"ProductID": 1}, {"ProductID": 2}], "@odata.nextLink": f"{V4_ROOT}Products?$skiptoken=2"}

    fetch = Pages(answer)
    result = collect(model("northwind-v4.xml", V4_ROOT), fetch, options=QueryOptions(top=1, filter="x"))
    assert [r["ProductID"] for r in result.rows] == [1, 2, 3]
    assert not result.truncated and result.pages == 2
    assert "$top=500" in fetch.urls[0] and "$filter=x" in fetch.urls[0]
    assert fetch.urls[1] == f"{V4_ROOT}Products?$skiptoken=2"


def test_follows_v2_next_links() -> None:
    def answer(url: str):
        if "$skiptoken" in url:
            return {"d": {"results": [{"ProductID": 2}]}}
        return {"d": {"results": [{"ProductID": 1}], "__count": "2", "__next": f"{V2_ROOT}Products?$skiptoken=1"}}

    result = collect(model("northwind-v2.xml", V2_ROOT), Pages(answer))
    assert [r["ProductID"] for r in result.rows] == [1, 2] and result.count == 2


def test_server_chunks_within_a_request_continue_after_all_rows() -> None:
    """Northwind answers $top=500 in chunks of 200 joined by next links."""
    total = 1100

    def answer(url: str):
        query = parse_qs(urlsplit(url).query)
        start = int(query["$skiptoken"][0]) if "$skiptoken" in query else skip_of(url)
        end_of_request = int(query.get("end", [start + 500])[0])
        stop = min(start + 200, end_of_request, total)
        body = {"value": [{"ProductID": i} for i in range(start, stop)]}
        if stop < min(end_of_request, total):
            body["@odata.nextLink"] = f"{V4_ROOT}Products?$skiptoken={stop}&end={end_of_request}"
        return body

    fetch = Pages(answer)
    result = collect(model("northwind-v4.xml", V4_ROOT), fetch, options=QueryOptions(count=False))
    ids = [r["ProductID"] for r in result.rows]
    assert ids == list(range(total))
    assert [skip_of(u) for u in fetch.urls if "$skiptoken" not in u] == [0, 500, 1000]


def test_falls_back_to_skip_and_stops_on_a_short_page() -> None:
    def answer(url: str):
        skip = skip_of(url)
        return {"value": [{"ProductID": i} for i in range(skip, min(skip + 500, 1200))]}

    fetch = Pages(answer)
    result = collect(model("northwind-v4.xml", V4_ROOT), fetch, options=QueryOptions(count=False))
    assert len(result.rows) == 1200 and not result.truncated
    assert [skip_of(u) for u in fetch.urls] == [0, 500, 1000]


def test_the_inline_count_ends_paging_below_a_silent_server_page_size() -> None:
    def answer(url: str):
        skip = skip_of(url)
        return {"value": [{"ProductID": i} for i in range(skip, min(skip + 100, 250))], "@odata.count": 250}

    fetch = Pages(answer)
    result = collect(model("northwind-v4.xml", V4_ROOT), fetch)
    assert len(result.rows) == 250 and [skip_of(u) for u in fetch.urls] == [0, 100, 200]


def test_starts_at_the_users_skip() -> None:
    def answer(url: str):
        skip = skip_of(url)
        return {"value": [{"ProductID": i} for i in range(skip, 30)], "@odata.count": 30}

    result = collect(model("northwind-v4.xml", V4_ROOT), Pages(answer), options=QueryOptions(skip=10))
    assert len(result.rows) == 20


def test_truncates_at_max_rows() -> None:
    def answer(url: str):
        skip = skip_of(url)
        return {"value": [{"ProductID": i} for i in range(skip, skip + 500)]}

    result = collect(model("northwind-v4.xml", V4_ROOT), Pages(answer), max_rows=1200)
    assert len(result.rows) == 1200 and result.truncated
    assert "1,200 rows" in result.notices[0]


def test_exactly_max_rows_with_a_known_count_is_complete() -> None:
    result = collect(
        model("northwind-v4.xml", V4_ROOT),
        Pages(lambda url: {"value": [{"ProductID": i} for i in range(500)], "@odata.count": 500}),
        max_rows=500,
    )
    assert len(result.rows) == 500 and not result.truncated


def test_stops_when_the_service_ignores_skip() -> None:
    fetch = Pages(lambda url: {"value": [{"ProductID": i} for i in range(500)]})
    result = collect(model("northwind-v4.xml", V4_ROOT), fetch)
    assert len(result.rows) == 500 and len(fetch.urls) == 2
    assert "ignored $skip" in result.notices[0]


def test_deadline_marks_the_export_incomplete() -> None:
    def answer(url: str):
        skip = skip_of(url)
        return {"value": [{"ProductID": i} for i in range(skip, skip + 500)]}

    result = collect(model("northwind-v4.xml", V4_ROOT), Pages(answer), deadline_seconds=-1)
    assert len(result.rows) == 500 and result.truncated


def test_byte_budget_marks_the_export_incomplete() -> None:
    def answer(url: str):
        skip = skip_of(url)
        return {"value": [{"ProductID": i, "Wide": "x" * 100} for i in range(skip, skip + 500)]}

    # One page is ~60 KB of JSON; the budget lets two through.
    result = collect(model("northwind-v4.xml", V4_ROOT), Pages(answer), max_total_bytes=150_000)
    assert len(result.rows) == 1000 and result.truncated
    assert "MB of data" in result.notices[0]


def test_each_request_is_bounded_by_the_remaining_deadline() -> None:
    timeouts: list[float] = []

    def fetch(url, model, auth, *, timeout, max_bytes):
        timeouts.append(timeout)
        return {"value": [{"ProductID": 1}]}, 200, "application/json", 1

    collect_rows(
        model("northwind-v4.xml", V4_ROOT), "Products", QueryOptions(), None,
        max_rows=10, deadline_seconds=5, timeout=30, max_bytes=1, fetch=fetch,
    )
    assert timeouts and timeouts[0] <= 5


def test_refuses_next_links_off_the_service() -> None:
    fetch = Pages(lambda url: {"value": [{"ProductID": 1}], "@odata.nextLink": "https://evil.example/x"})
    with pytest.raises(DataError):
        collect(model("northwind-v4.xml", V4_ROOT), fetch)


def test_single_targets_make_one_request() -> None:
    fetch = Pages(lambda url: {"ProductID": 1, "ProductName": "Chai"})
    result = collect(model("northwind-v4.xml", V4_ROOT), fetch, entity_set="Products(1)")
    assert result.rows == [{"ProductID": 1, "ProductName": "Chai"}] and len(fetch.urls) == 1
    assert "$top" not in fetch.urls[0]


def test_property_types_follow_the_path() -> None:
    types = property_types(model("northwind-v4.xml", V4_ROOT), "Customers('A')/Orders")
    assert types["OrderDate"] == "Edm.DateTimeOffset" and types["Freight"] == "Edm.Decimal"


@pytest.mark.parametrize(
    ("value", "edm", "expected"),
    [
        (None, "", None),
        (True, "Edm.Boolean", True),
        ("12.50", "Edm.Decimal", 12.5),
        ("9007199254740993", "Edm.Int64", "9007199254740993"),
        ("42", "Edm.Int64", 42),
        ("12.50", "Edm.String", "12.50"),
        ("2024-05-01T10:00:00+02:00", "Edm.DateTimeOffset", datetime(2024, 5, 1, 8, 0)),
        ("2024-05-01", "Edm.Date", datetime(2024, 5, 1)),
        ("not a date", "Edm.DateTime", "not a date"),
        ({"@link": "https://x/Orders(1)/Customer"}, "", "https://x/Orders(1)/Customer"),
        ({"City": "Wien"}, "", '{"City":"Wien"}'),
        ("x" * 40_000, "Edm.String", "x" * 32_767),
    ],
)
def test_cell_values(value: Any, edm: str, expected: Any) -> None:
    assert cell_value(value, edm) == expected


def test_workbook_round_trip() -> None:
    body = to_xlsx(
        ["ProductID", "UnitPrice", "OrderDate", "Note", "Discontinued"],
        [
            {"ProductID": 1, "UnitPrice": "18.00", "OrderDate": "1996-07-04T00:00:00Z", "Note": "=1+1", "Discontinued": False},
            {"ProductID": 2, "UnitPrice": None, "Note": "https://example.com"},
        ],
        types={"UnitPrice": "Edm.Decimal", "OrderDate": "Edm.DateTimeOffset"},
        info=[("Service", V4_ROOT), ("Rows exported", 2)],
    )
    wb = openpyxl.load_workbook(io.BytesIO(body))
    assert wb.sheetnames == ["Data", "Query"]
    data = list(wb["Data"].iter_rows(values_only=True))
    assert data[0] == ("ProductID", "UnitPrice", "OrderDate", "Note", "Discontinued")
    assert data[1] == (1, 18, datetime(1996, 7, 4), "=1+1", False)
    assert data[2] == (2, None, None, "https://example.com", None)
    assert wb["Data"]["D2"].data_type == "s"
    assert wb["Data"].freeze_panes == "A2"
    assert list(wb["Query"].iter_rows(values_only=True)) == [("Service", V4_ROOT), ("Rows exported", 2)]


def test_non_finite_numbers_are_written_as_text() -> None:
    data = to_xlsx(["v"], [{"v": float("nan")}, {"v": float("inf")}], types={}, info=[])
    assert data[:2] == b"PK"


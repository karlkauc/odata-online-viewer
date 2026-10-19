"""HTTP API: metadata loaders, source/export, samples, data proxy — with the network mocked."""

from __future__ import annotations

import io
import json
from collections.abc import Iterator
from dataclasses import asdict
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

import app.api.data as data_module
import app.api.metadata as metadata_module
from app.cache import metadata_cache
from app.main import app
from app.parser.security import FetchedResource, FetchHttpError, SecurityError
from app.rate_limit import limiter
from app.usage.context import UsageTracker
from app.usage.events import UsageEvent
from app.usage.recorder import UsageRecorder

SAMPLES = Path(__file__).resolve().parent.parent / "app" / "odata" / "samples"
NORTHWIND_V4 = (SAMPLES / "northwind-v4.xml").read_bytes()
NORTHWIND_V2 = (SAMPLES / "northwind-v2.xml").read_bytes()


class ListRecorder(UsageRecorder):
    def __init__(self) -> None:
        super().__init__("postgresql://u@h/db", "")
        self.events: list[UsageEvent] = []

    def record(self, event: UsageEvent) -> bool:
        self.events.append(event)
        return True

    async def start(self) -> None:
        return None

    async def stop(self) -> None:
        return None

    async def drain(self, timeout: float = 2.0) -> None:
        return None


@pytest.fixture
def recorder() -> ListRecorder:
    return ListRecorder()


@pytest.fixture
def client(recorder: ListRecorder) -> Iterator[TestClient]:
    limiter.reset()
    metadata_cache.clear()
    # No context manager: the lifespan would replace the tracker with the (disabled) production one.
    app.state.usage = UsageTracker(recorder, None, "secret")
    yield TestClient(app)
    app.state.usage = None


class FakeFetch:
    """Stand-in for security.fetch_url: records calls, returns canned answers."""

    def __init__(self) -> None:
        self.calls: list[tuple[str, dict[str, str]]] = []
        self.responses: dict[str, FetchedResource | Exception] = {}

    def __call__(self, url: str, *, headers=None, timeout=None, max_bytes=None) -> FetchedResource:
        self.calls.append((url, dict(headers or {})))
        for prefix, answer in self.responses.items():
            if url.startswith(prefix):
                if isinstance(answer, Exception):
                    raise answer
                return answer
        raise AssertionError(f"unexpected fetch {url}")


@pytest.fixture
def fetch(monkeypatch: pytest.MonkeyPatch) -> FakeFetch:
    fake = FakeFetch()
    monkeypatch.setattr(metadata_module, "fetch_url", fake)
    monkeypatch.setattr(data_module, "fetch_json", _fetch_json_via(fake))
    return fake


def _fetch_json_via(fake: FakeFetch):
    from app.odata import data_client

    def fetch_json(url, model, auth, *, timeout, max_bytes):
        original = data_client.fetch_url
        data_client.fetch_url = fake  # type: ignore[assignment]
        try:
            return data_client.fetch_json(url, model, auth, timeout=timeout, max_bytes=max_bytes)
        finally:
            data_client.fetch_url = original

    return fetch_json


# --- Loaders ----------------------------------------------------------------


def test_text_load(client: TestClient, recorder: ListRecorder) -> None:
    r = client.post("/api/metadata/text", json={"content": NORTHWIND_V4.decode(), "filename": "dir/nw.xml"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["odata_version"] == "4.0"
    assert body["service_root"] is None
    assert body["stats"]["entity_type_count"] == 26
    assert len(body["metadata_id"]) == 32
    ev = recorder.events[-1]
    assert (ev.event_type, ev.source, ev.status, ev.schema_name) == ("metadata_load", "text", "ok", "nw.xml")
    assert ev.element_count == 26 and ev.odata_version == "4.0"
    assert ev.target_namespace == "NorthwindModel"


def test_upload_load_and_source_round_trip(client: TestClient) -> None:
    r = client.post("/api/metadata/upload", files={"file": ("a/b/nw.xml", NORTHWIND_V2, "application/xml")})
    assert r.status_code == 200, r.text
    metadata_id = r.json()["metadata_id"]
    assert r.json()["filename"] == "a/b/nw.xml"
    again = client.get(f"/api/metadata/{metadata_id}")
    assert again.status_code == 200 and again.json()["metadata_id"] == metadata_id
    source = client.get(f"/api/metadata/{metadata_id}/source")
    assert source.status_code == 200
    assert source.headers["content-type"].startswith("text/plain")
    assert "sandbox" in source.headers["content-security-policy"]
    # Served pretty-printed (the snapshot is one long line), same content.
    assert r.json()["source_reformatted"] is True
    assert source.text.count("\n") > 100 and 'EntityType Name="Order"' in source.text
    md = client.get(f"/api/metadata/{metadata_id}/export/markdown")
    assert md.status_code == 200 and "attachment" in md.headers["content-disposition"]
    assert md.text.startswith("# ")
    mm = client.get(f"/api/metadata/{metadata_id}/export/mermaid")
    assert mm.text.startswith("erDiagram")
    assert client.get(f"/api/metadata/{metadata_id}/export/nope").status_code == 404


def test_parse_error_is_400_with_event(client: TestClient, recorder: ListRecorder) -> None:
    r = client.post("/api/metadata/text", json={"content": "<root/>", "filename": "x.xml"})
    assert r.status_code == 400
    assert "root element is <root>" in r.json()["detail"]
    ev = recorder.events[-1]
    assert ev.status == "parse_error" and ev.status_code == 400


def test_markdown_fence_is_removed_with_a_notice(client: TestClient) -> None:
    fenced = "```xml\n" + NORTHWIND_V4.decode() + "\n```"
    r = client.post("/api/metadata/text", json={"content": fenced, "filename": "x.xml"})
    assert r.status_code == 200, r.text
    assert any("Markdown code fence" in n for n in r.json()["notices"])


def test_expired_metadata_is_404(client: TestClient) -> None:
    r = client.get("/api/metadata/0123456789abcdef0123456789abcdef")
    assert r.status_code == 404
    assert r.json()["detail"] == "Metadata not found or expired"


def test_url_load_derives_service_root_after_redirect(
    client: TestClient, fetch: FakeFetch, recorder: ListRecorder
) -> None:
    final = "https://services.odata.org/TripPinRESTierService/(S(k1))/$metadata"
    fetch.responses["https://services.odata.org/TripPinRESTierService/$metadata"] = FetchedResource(
        url=final, content=(SAMPLES / "trippin-v4.xml").read_bytes(), content_type="application/xml"
    )
    r = client.post("/api/metadata/url", json={"url": "services.odata.org/TripPinRESTierService"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["metadata_url"] == final
    assert body["service_root"] == "https://services.odata.org/TripPinRESTierService/(S(k1))/"
    assert body["auth_used"] is False
    url, headers = fetch.calls[0]
    assert url == "https://services.odata.org/TripPinRESTierService/$metadata"
    assert headers["Accept"].startswith("application/xml")
    assert "Authorization" not in headers
    ev = recorder.events[-1]
    assert (
        ev.source == "url" and ev.schema_name == "https://services.odata.org/TripPinRESTierService/$metadata"
    )


def test_url_load_with_credentials_sends_basic_auth_but_never_logs_it(
    client: TestClient, fetch: FakeFetch, recorder: ListRecorder
) -> None:
    fetch.responses["https://odata.status20.net/odata/$metadata"] = FetchedResource(
        url="https://odata.status20.net/odata/$metadata",
        content=(SAMPLES / "jdbcodata-demo.xml").read_bytes(),
        content_type="application/xml",
    )
    r = client.post(
        "/api/metadata/url",
        json={
            "url": "https://odata.status20.net/odata/",
            "auth": {"username": "excel_demo", "password": "s3cret"},
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["auth_used"] is True
    assert fetch.calls[0][1]["Authorization"] == "Basic ZXhjZWxfZGVtbzpzM2NyZXQ="
    dumped = json.dumps([asdict(e) for e in recorder.events]) + r.text
    assert "s3cret" not in dumped and "ZXhjZWxfZGVtbzpzM2NyZXQ" not in dumped


def test_url_load_with_credentials_refuses_a_redirect_to_another_host(
    client: TestClient, fetch: FakeFetch
) -> None:
    # The data explorer would send the credentials to the redirect target.
    fetch.responses["https://odata.status20.net/odata/$metadata"] = FetchedResource(
        url="https://evil.example/svc/$metadata",
        content=(SAMPLES / "jdbcodata-demo.xml").read_bytes(),
        content_type="application/xml",
    )
    r = client.post(
        "/api/metadata/url",
        json={"url": "https://odata.status20.net/odata/", "auth": {"username": "u", "password": "p"}},
    )
    assert r.status_code == 400
    assert "https://evil.example" in r.json()["detail"]


def test_url_load_401_is_reported_as_auth_required(
    client: TestClient, fetch: FakeFetch, recorder: ListRecorder
) -> None:
    fetch.responses["https://odata.status20.net/odata/$metadata"] = FetchHttpError(
        "x", 401, www_authenticate='Basic realm="Realm"'
    )
    r = client.post("/api/metadata/url", json={"url": "https://odata.status20.net/odata/"})
    assert r.status_code == 401
    assert "requires authentication" in r.json()["detail"]
    ev = recorder.events[-1]
    assert ev.status == "rejected" and ev.status_code == 401


def test_url_load_security_error_is_400(client: TestClient, fetch: FakeFetch) -> None:
    fetch.responses["https://h.example/odata/$metadata"] = SecurityError(
        "host 'h.example' resolves to a private/loopback address"
    )
    r = client.post("/api/metadata/url", json={"url": "https://h.example/odata/"})
    assert r.status_code == 400 and "private" in r.json()["detail"]


def test_url_load_html_login_page_is_explained(client: TestClient, fetch: FakeFetch) -> None:
    fetch.responses["https://h.example/odata/$metadata"] = FetchedResource(
        url="https://h.example/odata/$metadata",
        content=b"<!doctype html><html><body>Sign in</body></html>",
        content_type="text/html",
    )
    r = client.post("/api/metadata/url", json={"url": "https://h.example/odata/"})
    assert r.status_code == 400 and "HTML page" in r.json()["detail"]


def test_url_of_an_entity_set_falls_back_to_the_service_root(client: TestClient, fetch: FakeFetch) -> None:
    fetch.responses["https://odata.status20.net/odata/Organisation/$metadata"] = FetchHttpError("x", 400)
    fetch.responses["https://odata.status20.net/odata/$metadata"] = FetchedResource(
        url="https://odata.status20.net/odata/$metadata",
        content=(SAMPLES / "jdbcodata-demo.xml").read_bytes(),
        content_type="application/xml",
    )
    r = client.post(
        "/api/metadata/url",
        json={
            "url": "https://odata.status20.net/odata/Organisation",
            "auth": {"username": "u", "password": "p"},
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["service_root"] == "https://odata.status20.net/odata/"
    assert any("service root https://odata.status20.net/odata/" in n for n in body["notices"])
    assert [c[0] for c in fetch.calls] == [
        "https://odata.status20.net/odata/Organisation/$metadata",
        "https://odata.status20.net/odata/$metadata",
    ]
    assert all(c[1].get("Authorization") for c in fetch.calls)


def test_parent_fallback_stops_at_auth_errors_and_reports_the_first_error(
    client: TestClient, fetch: FakeFetch
) -> None:
    fetch.responses["https://h.example/a/b/$metadata"] = FetchHttpError("x", 404)
    fetch.responses["https://h.example/a/$metadata"] = FetchHttpError("x", 401)
    r = client.post("/api/metadata/url", json={"url": "https://h.example/a/b"})
    assert r.status_code == 401 and len(fetch.calls) == 2
    fetch.calls.clear()
    fetch.responses.clear()
    fetch.responses["https://h.example/x/$metadata"] = FetchHttpError("x", 404)
    fetch.responses["https://h.example/$metadata"] = FetchedResource(
        url="https://h.example/$metadata",
        content=b"<!doctype html><html>home</html>",
        content_type="text/html",
    )
    r = client.post("/api/metadata/url", json={"url": "https://h.example/x/"})
    assert r.status_code == 400 and "HTTP 404" in r.json()["detail"]
    # an explicit $metadata URL is not second-guessed
    fetch.calls.clear()
    fetch.responses["https://h.example/x/$metadata"] = FetchHttpError("x", 400)
    r = client.post("/api/metadata/url", json={"url": "https://h.example/x/$metadata"})
    assert r.status_code == 400 and len(fetch.calls) == 1


# --- Samples ----------------------------------------------------------------


def test_samples_list_and_bundled_load(client: TestClient, fetch: FakeFetch) -> None:
    r = client.get("/api/samples")
    ids = [s["id"] for s in r.json()["samples"]]
    assert "northwind-v4" in ids and "trippin-v4" in ids
    assert all("file" not in s for s in r.json()["samples"])
    r = client.post("/api/samples/northwind-v4/load", params={"live": "false"})
    assert r.status_code == 200, r.text
    assert r.json()["source"] == "sample" and r.json()["service_root"] is None
    assert client.post("/api/samples/nope/load").status_code == 404


def test_sample_live_failure_falls_back_to_snapshot(client: TestClient, fetch: FakeFetch) -> None:
    fetch.responses["https://services.odata.org/V4/Northwind/Northwind.svc/$metadata"] = FetchHttpError(
        "down", 503
    )
    r = client.post("/api/samples/northwind-v4/load")
    assert r.status_code == 200, r.text
    assert any("not reachable" in n for n in r.json()["notices"])
    assert r.json()["service_root"] is None


def test_sample_live_load_keeps_service_root(client: TestClient, fetch: FakeFetch) -> None:
    url = "https://services.odata.org/V4/Northwind/Northwind.svc/$metadata"
    fetch.responses[url] = FetchedResource(url=url, content=NORTHWIND_V4, content_type="application/xml")
    r = client.post("/api/samples/northwind-v4/load")
    assert r.status_code == 200
    assert r.json()["service_root"] == "https://services.odata.org/V4/Northwind/Northwind.svc/"
    assert r.json()["source"] == "sample"


# --- Data proxy -------------------------------------------------------------


def _loaded_v4(client: TestClient, fetch: FakeFetch) -> str:
    url = "https://services.odata.org/V4/Northwind/Northwind.svc/$metadata"
    fetch.responses[url] = FetchedResource(url=url, content=NORTHWIND_V4, content_type="application/xml")
    return client.post(
        "/api/metadata/url", json={"url": "https://services.odata.org/V4/Northwind/Northwind.svc/"}
    ).json()["metadata_id"]


def test_data_query_v4(client: TestClient, fetch: FakeFetch, recorder: ListRecorder) -> None:
    metadata_id = _loaded_v4(client, fetch)
    fetch.responses["https://services.odata.org/V4/Northwind/Northwind.svc/Products?"] = FetchedResource(
        url="…",
        content=json.dumps(
            {
                "@odata.count": 2,
                "value": [{"ProductID": 1, "ProductName": "Chai"}, {"ProductID": 2, "ProductName": "Chang"}],
            }
        ).encode(),
        content_type="application/json;odata.metadata=minimal",
    )
    r = client.post(
        "/api/data/query",
        json={
            "metadata_id": metadata_id,
            "entity_set": "Products",
            "options": {"top": 2, "select": ["ProductID", "ProductName"]},
            "auth": {"username": "u", "password": "p"},
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["columns"] == ["ProductID", "ProductName"]
    assert body["count"] == 2 and len(body["rows"]) == 2
    assert body["request_url"].startswith(
        "https://services.odata.org/V4/Northwind/Northwind.svc/Products?$select=ProductID,ProductName&$top=2"
    )
    url, headers = fetch.calls[-1]
    assert headers["Authorization"] == "Basic dTpw"
    assert headers["Accept"].startswith("application/json")
    ev = recorder.events[-1]
    assert (ev.event_type, ev.source, ev.status, ev.element_count) == ("data_query", "query", "ok", 2)
    assert ev.schema_name == "https://services.odata.org/V4/Northwind/Northwind.svc/"
    assert ev.target_namespace == "Products"
    assert "dTpw" not in json.dumps([asdict(e) for e in recorder.events])


def test_data_query_collection_property_is_not_paged(client: TestClient, fetch: FakeFetch) -> None:
    root = "https://services.odata.org/TripPinRESTierService/"
    fetch.responses[root + "$metadata"] = FetchedResource(
        url=root + "$metadata", content=(SAMPLES / "trippin-v4.xml").read_bytes(), content_type="application/xml"
    )
    metadata_id = client.post("/api/metadata/url", json={"url": root}).json()["metadata_id"]
    fetch.responses[root + "People('russellwhyte')/Features"] = FetchedResource(
        url="…",
        content=json.dumps({"value": ["Feature1", "Feature4"]}).encode(),
        content_type="application/json;odata.metadata=minimal",
    )
    r = client.post(
        "/api/data/query",
        json={"metadata_id": metadata_id, "entity_set": "People('russellwhyte')/Features", "options": {}},
    )
    assert r.status_code == 200, r.text
    assert r.json()["count"] == 2 and len(r.json()["rows"]) == 2
    # TripPin answers $count on a collection property with HTTP 500.
    assert fetch.calls[-1][0] == root + "People('russellwhyte')/Features"


def test_data_query_upstream_error(client: TestClient, fetch: FakeFetch, recorder: ListRecorder) -> None:
    metadata_id = _loaded_v4(client, fetch)
    fetch.responses["https://services.odata.org/V4/Northwind/Northwind.svc/Products?"] = FetchHttpError(
        "x", 400
    )
    r = client.post(
        "/api/data/query",
        json={"metadata_id": metadata_id, "entity_set": "Products", "options": {"filter": "bogus"}},
    )
    assert r.status_code == 502
    assert "HTTP 400" in r.json()["detail"]
    assert recorder.events[-1].status == "error" and recorder.events[-1].status_code == 400


def test_data_query_without_service_root(client: TestClient) -> None:
    metadata_id = client.post("/api/metadata/text", json={"content": NORTHWIND_V4.decode()}).json()[
        "metadata_id"
    ]
    r = client.post("/api/data/query", json={"metadata_id": metadata_id, "entity_set": "Products"})
    assert r.status_code == 400 and "pasted or uploaded" in r.json()["detail"]


def test_data_query_top_is_capped(client: TestClient, fetch: FakeFetch) -> None:
    metadata_id = _loaded_v4(client, fetch)
    r = client.post(
        "/api/data/query",
        json={"metadata_id": metadata_id, "entity_set": "Products", "options": {"top": 5000}},
    )
    assert r.status_code == 422


def test_data_raw(client: TestClient, fetch: FakeFetch) -> None:
    metadata_id = _loaded_v4(client, fetch)
    fetch.responses["https://services.odata.org/V4/Northwind/Northwind.svc/Products(1)"] = FetchedResource(
        url="…", content=json.dumps({"ProductID": 1}).encode(), content_type="application/json"
    )
    r = client.post(
        "/api/data/raw", json={"metadata_id": metadata_id, "path_and_query": "Products(1)?$select=ProductID"}
    )
    assert r.status_code == 200, r.text
    assert r.json()["body"] == {"ProductID": 1}
    assert (
        r.json()["request_url"]
        == "https://services.odata.org/V4/Northwind/Northwind.svc/Products(1)?$select=ProductID"
    )
    bad = client.post("/api/data/raw", json={"metadata_id": metadata_id, "path_and_query": "../admin"})
    assert bad.status_code == 400


def test_data_next_link_must_be_on_the_service(client: TestClient, fetch: FakeFetch) -> None:
    metadata_id = _loaded_v4(client, fetch)
    r = client.post(
        "/api/data/query",
        json={
            "metadata_id": metadata_id,
            "entity_set": "Products",
            "next_link": "https://evil.example/Products",
        },
    )
    assert r.status_code == 400 and "outside" in r.json()["detail"]


# --- Excel export -----------------------------------------------------------


def test_data_export_page(client: TestClient, fetch: FakeFetch, recorder: ListRecorder) -> None:
    metadata_id = _loaded_v4(client, fetch)
    calls = len(fetch.calls)
    r = client.post(
        "/api/data/export",
        json={
            "metadata_id": metadata_id,
            "entity_set": "Products",
            "scope": "page",
            "columns": ["ProductID", "UnitPrice"],
            "rows": [{"ProductID": 1, "UnitPrice": "18.00"}],
        },
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["x-export-rows"] == "1" and r.headers["x-export-truncated"] == "0"
    assert 'filename="Products-' in r.headers["content-disposition"]
    assert len(fetch.calls) == calls, "a page export must not ask the service again"
    rows = list(openpyxl.load_workbook(io.BytesIO(r.content))["Data"].iter_rows(values_only=True))
    assert rows == [("ProductID", "UnitPrice"), (1, 18)]
    ev = recorder.events[-1]
    assert (ev.event_type, ev.source, ev.element_count, ev.target_namespace) == ("export", "xlsx_page", 1, "Products")


def test_data_export_page_needs_rows(client: TestClient, fetch: FakeFetch) -> None:
    metadata_id = _loaded_v4(client, fetch)
    r = client.post("/api/data/export", json={"metadata_id": metadata_id, "entity_set": "Products", "scope": "page"})
    assert r.status_code == 422


def test_data_export_all(client: TestClient, fetch: FakeFetch, recorder: ListRecorder) -> None:
    metadata_id = _loaded_v4(client, fetch)
    root = "https://services.odata.org/V4/Northwind/Northwind.svc/"
    fetch.responses[f"{root}Products?$skiptoken"] = FetchedResource(
        url="…", content=json.dumps({"value": [{"ProductID": 2}]}).encode(), content_type="application/json"
    )
    fetch.responses[f"{root}Products?"] = FetchedResource(
        url="…",
        content=json.dumps({"value": [{"ProductID": 1}], "@odata.nextLink": f"{root}Products?$skiptoken=1"}).encode(),
        content_type="application/json",
    )
    r = client.post(
        "/api/data/export",
        json={
            "metadata_id": metadata_id,
            "entity_set": "Products",
            "scope": "all",
            "options": {"select": ["ProductID"]},
            "auth": {"username": "u", "password": "p"},
        },
    )
    assert r.status_code == 200, r.text
    assert r.headers["x-export-rows"] == "2"
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert list(wb["Data"].iter_rows(values_only=True)) == [("ProductID",), (1,), (2,)]
    query = dict(wb["Query"].iter_rows(values_only=True))
    assert query["$select"] == "ProductID" and query["Complete"] == "yes"
    assert all(h["Authorization"] == "Basic dTpw" for _, h in fetch.calls[-2:])
    assert "dTpw" not in json.dumps([asdict(e) for e in recorder.events])
    assert recorder.events[-1].source == "xlsx_all"


def test_data_export_all_upstream_error(client: TestClient, fetch: FakeFetch, recorder: ListRecorder) -> None:
    metadata_id = _loaded_v4(client, fetch)
    fetch.responses["https://services.odata.org/V4/Northwind/Northwind.svc/Products?"] = FetchHttpError("x", 401)
    r = client.post("/api/data/export", json={"metadata_id": metadata_id, "entity_set": "Products", "scope": "all"})
    assert r.status_code == 502 and "HTTP 401" in r.json()["detail"]
    assert recorder.events[-1].status == "error" and recorder.events[-1].source == "xlsx_all"


def test_data_export_unknown_metadata(client: TestClient) -> None:
    r = client.post("/api/data/export", json={"metadata_id": "0" * 32, "entity_set": "Products", "scope": "all"})
    assert r.status_code == 404
